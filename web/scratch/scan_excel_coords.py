import sys
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "database"))
from coord_parser import parse_and_validate_coordinates

wb = openpyxl.load_workbook('database/imports/Cơ sở dữ liệu xd bản đồ MSVT 13.10.xlsx', data_only=True)
ws = wb['Sheet1']

for row_idx in range(12, 63):
    coords = ws.cell(row=row_idx, column=8).value
    code = ws.cell(row=row_idx, column=2).value
    if not coords:
        continue
    res = parse_and_validate_coordinates(coords)
    if res['total_groups'] > 1 or res['total_points'] > 5:
        print(f"Row {row_idx} ({code}): {res['total_groups']} groups, {res['total_points']} pts, all_valid={res['all_groups_valid']}")
        for g in res['groups']:
            print(f"   -> {g['name']}: {len(g['points'])} pts, valid={g['is_valid']}, geom={g['geom_type']}, reason={g['reason']}")
